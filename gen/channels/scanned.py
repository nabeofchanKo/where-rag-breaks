"""`scanned` チャネル — 答えがテキスト層のない画像PDFの中にある。

SPEC §3-1 #6 / §3-4。PyMuPDF で PDF を作り、各ページをラスタライズして
画像だけの PDF に組み直す。

なぜ古典的RAGが落ちるのか:
    文字が画素になっているので、pypdf も pdfplumber も**何も返さない**。
    これは抽出器の穴ではなく、テキスト層が存在しないという構造である。
    実際 `tests/test_channels_p1b.py` で「抽出結果が空であること」を固定している。

難易度 = 罠の機構:
    1 … **コントロール**。通常のテキストPDF（ラスタライズしない）。
    2 … 画像だけのPDF。
    3 … 画像だけのPDF。さらに送付状の docx に「前回ロットの判定は〜」という
        文を置く。PDF が読めないと、読める唯一の文書であるこの送付状から
        答えを拾おうとする。
"""

from __future__ import annotations

import random
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Pt

from gen.common import Item, normalize_artifact, normalize_pdf
from gen.imaging import save_scanned_pdf
from gen.locales import Locale

CHANNEL = "scanned"
SUBDIR = "inspection_reports"

# (difficulty, ラスタライズするか, 送付状に囮を置くか)
QUESTION_SPECS: tuple[tuple[int, bool, bool], ...] = (
    (1, False, False),
    (2, True, False),
    (3, True, True),
)


def _save_text_pdf(path: Path, lines: list[str]) -> None:
    """ラスタライズしない通常のテキストPDF（コントロール段用）。"""
    doc = pymupdf.open()
    page = doc.new_page()
    y = 80.0
    for line in lines:
        page.insert_text((64, y), line, fontsize=11, fontname="japan")
        y += 22.0
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path, deflate=True)
    doc.close()
    normalize_pdf(path)


def _build(
    rng: random.Random, outdir: Path, loc: Locale, code: str, rasterize: bool, with_decoy: bool
) -> tuple[str, str | None, list[str]]:
    judges = ("合格", "条件付合格") if loc.code == "ja" else ("Pass", "Conditional pass")
    lot = f"LOT-{rng.randrange(100000, 999999)}"
    judge = rng.choice(judges)

    lines = [
        loc.s["sc_title"],
        loc.fmt("sc_line_code", code=code),
        loc.fmt("sc_line_site", site=rng.choice(loc.sites)),
        loc.fmt("sc_line_item", item=rng.choice(loc.items)),
        loc.fmt("sc_line_lot", lot=lot),
        loc.fmt("sc_line_value", value=f"{rng.uniform(9.0, 11.0):.2f}"),
        loc.fmt("sc_line_judge", judge=judge),
    ]

    pdf_rel = f"{SUBDIR}/{code}_report.pdf"
    if rasterize:
        save_scanned_pdf(outdir / pdf_rel, [lines])
    else:
        _save_text_pdf(outdir / pdf_rel, lines)
    files = [pdf_rel]

    decoy: str | None = None
    if with_decoy:
        # 読める唯一の文書。ここに載っているのは**前回ロット**の番号であって
        # 設問が訊いているロット番号ではない。
        decoy = f"LOT-{rng.randrange(100000, 999999)}"
        if decoy == lot:
            decoy = None
        doc = Document()
        font = doc.styles["Normal"].font
        font.name = "Yu Gothic" if loc.code == "ja" else "Calibri"
        font.size = Pt(10.5)
        doc.add_heading(loc.fmt("sc_cover_title", code=code), level=1)
        doc.add_paragraph(loc.s["sc_cover_body"])
        if decoy:
            doc.add_paragraph(loc.fmt("sc_cover_decoy", lot=decoy, judge=judge))
        cover_rel = f"{SUBDIR}/{code}_cover.docx"
        path = outdir / cover_rel
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        normalize_artifact(path)
        files.append(cover_rel)

    return lot, decoy, files


def generate(rng: random.Random, outdir: Path, n_questions: int, locale: Locale) -> list[Item]:
    codes = [f"SC-{n:04d}" for n in rng.sample(range(1000, 5000), n_questions)]

    items: list[Item] = []
    for i, code in enumerate(codes):
        difficulty, rasterize, with_decoy = QUESTION_SPECS[i % len(QUESTION_SPECS)]
        lot, decoy, files = _build(rng, outdir, locale, code, rasterize, with_decoy)
        items.append(
            Item(
                qid=f"{CHANNEL}-{i + 1:02d}",
                channel=CHANNEL,
                question=locale.fmt("q_scanned", code=code),
                answer=lot,
                answer_type="identifier",
                source_files=files,
                decoys=[decoy] if decoy else [],
                difficulty=difficulty,
                locale=locale.code,
            )
        )
    return items


def generate_fillers(rng: random.Random, outdir: Path, count: int, locale: Locale) -> list[str]:
    written: list[str] = []
    for n in rng.sample(range(5000, 9999), count):
        _, rasterize, with_decoy = QUESTION_SPECS[rng.randrange(len(QUESTION_SPECS))]
        _, _, files = _build(rng, outdir, locale, f"SC-{n:04d}", rasterize, with_decoy)
        written.extend(files)
    return written
