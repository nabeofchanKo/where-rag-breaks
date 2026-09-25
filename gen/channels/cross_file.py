"""`cross_file` チャネル — 答えが N ファイルにまたがる集計にしかない。

SPEC §3-1 #9。案件フォルダそれぞれに 1 つずつ数値を置き、合計を問う。

なぜ古典的RAGが落ちるのか:
    top-k は **k 個のチャンクしか載せられない**。10 ファイルに散らばった数値を
    全部集めるには k をそれ以上に上げる必要があり、しかも無関係なチャンクと
    枠を奪い合う。検索が失敗しているのではなく、窓が足りない。
    P0 の `formula` 難易度2 と同じ機構を、1 ファイル内ではなく
    **ファイル間**で起こす。

難易度 = 罠の機構:
    1 … **コントロール**。3 ファイルだけ。k=4 でも窓に収まる。
    2 … 10 ファイル。
    3 … 10 ファイル。さらに**陳腐化した集計表**を同じフォルダに置く。
        1 ファイルで答えが手に入るように見えるが、その値は古い。
        ``decoys`` に登録してあるので「囮を掴んだ割合」を測れる。
"""

from __future__ import annotations

import random
from pathlib import Path

from docx import Document
from docx.shared import Pt

from gen.common import Item, normalize_artifact, number_aliases
from gen.locales import Locale

CHANNEL = "cross_file"
SUBDIR = "cases"

# (difficulty, ファイル数, 陳腐化した集計表を置くか)
QUESTION_SPECS: tuple[tuple[int, int, bool], ...] = (
    (1, 3, False),
    (2, 10, False),
    (3, 10, True),
)


def _write_doc(path: Path, loc: Locale, title: str, body: str) -> None:
    doc = Document()
    font = doc.styles["Normal"].font
    font.name = "Yu Gothic" if loc.code == "ja" else "Calibri"
    font.size = Pt(10.5)
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    normalize_artifact(path)


def _build(
    rng: random.Random, outdir: Path, loc: Locale, series: str, n_files: int, stale: bool
) -> tuple[int, int | None, list[str], str, str]:
    amounts = [rng.randrange(120, 4800) * 1000 for _ in range(n_files)]
    total = sum(amounts)

    written: list[str] = []
    codes: list[str] = []
    for index, amount in enumerate(amounts, start=1):
        code = f"{series}-{index:02d}"
        codes.append(code)
        rel = f"{SUBDIR}/{series.lower()}/{code}_record.docx"
        _write_doc(
            outdir / rel,
            loc,
            loc.fmt("cf_doc_title", code=code),
            loc.fmt(
                "cf_body",
                code=code,
                amount=f"{amount:,}",
                department=rng.choice(loc.departments),
            ),
        )
        written.append(rel)

    stale_total: int | None = None
    if stale:
        # 1 件ぶん取りこぼした古い集計。**1 ファイルで答えが手に入るように見える**
        dropped = rng.randrange(n_files)
        stale_total = total - amounts[dropped]
        rel = f"{SUBDIR}/{series.lower()}/{series}_summary.docx"
        _write_doc(
            outdir / rel,
            loc,
            loc.fmt("cf_summary_title", series=series),
            loc.fmt("cf_summary_body", amount=f"{stale_total:,}", asof="2027-04-30"),
        )
        written.append(rel)

    return total, stale_total, written, codes[0], codes[-1]


def _unit(loc: Locale) -> str:
    return "円" if loc.code == "ja" else ""


def generate(rng: random.Random, outdir: Path, n_questions: int, locale: Locale) -> list[Item]:
    series_ids = [f"CF-{n:04d}" for n in rng.sample(range(1000, 5000), n_questions)]

    items: list[Item] = []
    for i, series in enumerate(series_ids):
        difficulty, n_files, stale = QUESTION_SPECS[i % len(QUESTION_SPECS)]
        total, stale_total, files, first, last = _build(
            rng, outdir, locale, series, n_files, stale
        )
        items.append(
            Item(
                qid=f"{CHANNEL}-{i + 1:02d}",
                channel=CHANNEL,
                question=locale.fmt("q_cross_file", series=series, first=first, last=last),
                answer=str(total),
                answer_type="number",
                answer_aliases=number_aliases(total, _unit(locale))[1:],
                source_files=files,
                decoys=[str(stale_total)] if stale_total is not None else [],
                difficulty=difficulty,
                locale=locale.code,
            )
        )
    return items


def generate_fillers(rng: random.Random, outdir: Path, count: int, locale: Locale) -> list[str]:
    """フィラーは 1 件 = 1 シリーズだと多すぎるので、単発の案件台帳だけ置く。"""
    written: list[str] = []
    for n in rng.sample(range(5000, 9999), count):
        code = f"CF-{n:04d}-01"
        rel = f"{SUBDIR}/misc/{code}_record.docx"
        _write_doc(
            outdir / rel,
            locale,
            locale.fmt("cf_doc_title", code=code),
            locale.fmt(
                "cf_body",
                code=code,
                amount=f"{rng.randrange(120, 4800) * 1000:,}",
                department=rng.choice(locale.departments),
            ),
        )
        written.append(rel)
    return written
