"""`locked` チャネル — 答えがパスワード付きファイルの中にある。

SPEC §3-1 #11 / §3-5。**パスワードそのものは書かず、パスワード規則を別文書に
書く**のがこのチャネルの肝。規則と材料（拠点コード・確定日）を別々の文書から
集めて組み立てないと開けられない。

構成:
    ``{code}_settlement.zip``  AES 暗号化。中に精算確定額（= 答え）
    ``{code}_handling.docx``   パスワード規則「拠点コード + ハイフン + 確定日」
    ``{code}_master.docx``     拠点コードと確定日（規則に入れる材料）

なぜ古典的RAGが落ちるのか:
    暗号化された zip はテキスト抽出の対象にならない。仮に規則と材料を
    両方引けても、**ファイルを開く手段が無い**。

★ 暗号化と決定性の両立（実装上の注意）:
    AES はエントリごとにランダムな salt を使うので、素直に作ると同 seed でも
    バイト一致しない。``gen/common.py`` の ``deterministic_aes_random`` で
    salt の生成源を固定している。**合成データであり、パスワードは別文書に
    公開されている前提のコーパス**なので、暗号強度は論点ではない。
    詳細は docs/environment-notes.md。

難易度 = 罠の機構:
    1 … **コントロール**。同じ金額が送付状にも書いてある。
    2 … 暗号化ファイルの中にしかない。
    3 … 同上。送付状に「速報値」として**別の金額**が書いてある。
"""

from __future__ import annotations

import datetime as dt
import random
from pathlib import Path

from docx import Document
from docx.shared import Pt

from gen.common import Item, deterministic_aes_zip, normalize_artifact, number_aliases
from gen.locales import Locale

CHANNEL = "locked"
SUBDIR = "settlements"

SITE_CODES = ("YKH", "KWS", "NGY", "OSK", "SDI", "FKO", "KYM", "KOB")

# (difficulty, 送付状に答えを出すか, 送付状に速報値の囮を置くか)
QUESTION_SPECS: tuple[tuple[int, bool, bool], ...] = (
    (1, True, False),
    (2, False, False),
    (3, False, True),
)


def _write_doc(path: Path, loc: Locale, title: str, lines: list[str]) -> None:
    doc = Document()
    font = doc.styles["Normal"].font
    font.name = "Yu Gothic" if loc.code == "ja" else "Calibri"
    font.size = Pt(10.5)
    doc.add_heading(title, level=1)
    for line in lines:
        doc.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    normalize_artifact(path)


def _build(
    rng: random.Random, outdir: Path, loc: Locale, code: str, show_amount: bool, with_decoy: bool
) -> tuple[int, int | None, list[str]]:
    amount = rng.randrange(800, 9800) * 10000
    site_code = rng.choice(SITE_CODES)
    date = dt.date(2027, rng.randrange(1, 13), rng.randrange(1, 28))
    password = f"{site_code}-{date.strftime('%Y%m%d')}"

    files: list[str] = []

    # ① 暗号化アーカイブ（答えはこの中だけ）
    zip_rel = f"{SUBDIR}/{code}_settlement.zip"
    deterministic_aes_zip(
        outdir / zip_rel,
        password=password,
        entries={
            loc.s["lk_archive_name"]: loc.fmt(
                "lk_archive_body", code=code, amount=f"{amount:,}", date=date.isoformat()
            )
        },
    )
    files.append(zip_rel)

    # ② パスワード規則（規則だけ。パスワードそのものは書かない）
    rule_rel = f"{SUBDIR}/{code}_handling.docx"
    _write_doc(
        outdir / rule_rel, loc, loc.fmt("lk_rule_title", code=code), [loc.s["lk_rule_body"]]
    )
    files.append(rule_rel)

    # ③ 規則に入れる材料（拠点コードと確定日）
    master_rel = f"{SUBDIR}/{code}_master.docx"
    _write_doc(
        outdir / master_rel,
        loc,
        loc.fmt("lk_ref_title", code=code),
        [
            loc.fmt(
                "lk_ref_body",
                site_code=site_code,
                date=date.isoformat(),
                department=rng.choice(loc.departments),
            )
        ],
    )
    files.append(master_rel)

    # ④ 送付状
    decoy: int | None = None
    cover_lines = [loc.s["lk_cover_body"]]
    if show_amount:
        cover_lines.append(loc.fmt("lk_archive_body", code=code, amount=f"{amount:,}",
                                   date=date.isoformat()))
    if with_decoy:
        decoy = rng.randrange(800, 9800) * 10000
        if decoy == amount:
            decoy = None
        else:
            cover_lines.append(loc.fmt("lk_cover_decoy", amount=f"{decoy:,}"))
    cover_rel = f"{SUBDIR}/{code}_cover.docx"
    _write_doc(outdir / cover_rel, loc, loc.fmt("lk_cover_title", code=code), cover_lines)
    files.append(cover_rel)

    return amount, decoy, files


def _unit(loc: Locale) -> str:
    return "円" if loc.code == "ja" else ""


def generate(rng: random.Random, outdir: Path, n_questions: int, locale: Locale) -> list[Item]:
    codes = [f"ST-{n:04d}" for n in rng.sample(range(1000, 5000), n_questions)]

    items: list[Item] = []
    for i, code in enumerate(codes):
        difficulty, show_amount, with_decoy = QUESTION_SPECS[i % len(QUESTION_SPECS)]
        amount, decoy, files = _build(rng, outdir, locale, code, show_amount, with_decoy)
        items.append(
            Item(
                qid=f"{CHANNEL}-{i + 1:02d}",
                channel=CHANNEL,
                question=locale.fmt("q_locked", code=code),
                answer=str(amount),
                answer_type="number",
                answer_aliases=number_aliases(amount, _unit(locale))[1:],
                source_files=files,
                decoys=[str(decoy)] if decoy is not None else [],
                difficulty=difficulty,
                locale=locale.code,
            )
        )
    return items


def generate_fillers(rng: random.Random, outdir: Path, count: int, locale: Locale) -> list[str]:
    written: list[str] = []
    for n in rng.sample(range(5000, 9999), count):
        _, show_amount, with_decoy = QUESTION_SPECS[rng.randrange(len(QUESTION_SPECS))]
        _, _, files = _build(rng, outdir, locale, f"ST-{n:04d}", show_amount, with_decoy)
        written.extend(files)
    return written
