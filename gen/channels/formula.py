"""`formula` チャネル — 答えが数式とその参照範囲にある。

SPEC §3-1 #3 / §3-5。**合計値はファイルのどこにも書かない。**

なぜこれが古典的RAGを壊すのか:
    openpyxl が書いた数式セルには**キャッシュ済み計算値が存在しない**。
    ``data_only=True`` で読むと ``None`` が返る。つまり抽出テキストに合計値は
    一切現れない。さらに金額列そのものも ``=数量*単価`` の数式なので、
    個々の金額すら現れない。抽出できるのは数量・単価・区分と、数式の文字列だけ。

    これは実装の手抜きではなく**構造**である（SPEC §11: 抽出器に意図的な穴を
    作らない）。Arm A の抽出器は値と数式文字列の両方を出す。そのうえで、
    行の対応関係を失ったテキストから積・条件付き合計・連鎖計算を復元できるか
    が問われる。

難易度:
    1 … 単純合計（SUM）
    2 … 別シート参照 + 条件付き合計（SUMIF）
    3 … 連鎖（SUM → 税率が別シート → 税込）
"""

from __future__ import annotations

import random
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from gen.common import Item, normalize_artifact, number_aliases
from gen.locales import Locale

CHANNEL = "formula"
SUBDIR = "quotes"

HEADER_ROW = 3
FIRST_DATA_ROW = 4
TAX_RATE = Decimal("0.10")

# (質問キー, 答えの種類, difficulty)
QUESTION_SPECS: tuple[tuple[str, str, int], ...] = (
    ("q_subtotal", "subtotal", 1),
    ("q_approved", "approved", 2),
    ("q_total", "total", 3),
)


def _excel_round(value: Decimal) -> int:
    """Excel の ROUND(x, 0) と同じ丸め（0 から遠い方へ half up）。

    Python 組込みの ``round`` は銀行家丸めなので使えない。正解が 1 円ずれる。
    """
    return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _make_lines(rng: random.Random, loc: Locale) -> list[dict]:
    """明細行を決める。ここで確定した値がそのまま正解計算の入力になる。"""
    n = rng.randrange(8, 13)
    items = rng.sample(loc.items, n)
    approved, on_hold = loc.statuses
    lines = []
    for name in items:
        lines.append(
            {
                "item": name,
                "qty": rng.randrange(1, 25),
                "unit_price": rng.randrange(12, 900) * 1000,
                # 3 割前後を保留にする。全部承認済だと difficulty 2 が成立しない
                "status": on_hold if rng.random() < 0.35 else approved,
            }
        )
    # 承認済と保留が最低1件ずつ存在することを保証する
    if all(line["status"] == approved for line in lines):
        lines[rng.randrange(n)]["status"] = on_hold
    if all(line["status"] == on_hold for line in lines):
        lines[rng.randrange(n)]["status"] = approved
    return lines


def _compute(lines: list[dict], approved_label: str) -> dict[str, int]:
    """正解を **生成時に** 計算する。ファイルを読み直して作らない（SPEC §3-3）。"""
    subtotal = sum(line["qty"] * line["unit_price"] for line in lines)
    approved = sum(
        line["qty"] * line["unit_price"] for line in lines if line["status"] == approved_label
    )
    tax = _excel_round(Decimal(subtotal) * TAX_RATE)
    return {
        "subtotal": subtotal,
        "approved": approved,
        "tax": tax,
        "total": subtotal + tax,
    }


def _write_workbook(
    path: Path, loc: Locale, code: str, project: str, lines: list[dict]
) -> None:
    """明細 / サマリ / 設定 の3シートを書く。**計算結果はどこにも書かない。**"""
    wb = Workbook()
    detail = wb.active
    detail.title = loc.s["xl_sheet_detail"]
    summary = wb.create_sheet(loc.s["xl_sheet_summary"])
    config = wb.create_sheet(loc.s["xl_sheet_config"])

    # ── 明細シート ────────────────────────────────────────
    detail["A1"] = loc.fmt("xl_title", project=project, code=code)
    detail["A1"].font = Font(bold=True, size=13)
    detail["A2"] = loc.s["xl_note"]

    headers = ("xl_h_item", "xl_h_qty", "xl_h_unit_price", "xl_h_amount", "xl_h_status")
    for col, key in enumerate(headers, start=1):
        cell = detail.cell(row=HEADER_ROW, column=col, value=loc.s[key])
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for offset, line in enumerate(lines):
        r = FIRST_DATA_ROW + offset
        detail.cell(row=r, column=1, value=line["item"])
        detail.cell(row=r, column=2, value=line["qty"])
        detail.cell(row=r, column=3, value=line["unit_price"])
        # 金額そのものが数式。抽出テキストに個々の金額すら現れない。
        detail.cell(row=r, column=4, value=f"=B{r}*C{r}")
        detail.cell(row=r, column=5, value=line["status"])

    last_row = FIRST_DATA_ROW + len(lines) - 1
    for col, width in enumerate((26, 8, 12, 14, 12), start=1):
        detail.column_dimensions[get_column_letter(col)].width = width

    # ── 設定シート（税率は別シートに置く。difficulty 3 の連鎖の起点）──
    config["A1"] = loc.s["xl_sheet_config"]
    config["A1"].font = Font(bold=True)
    config["A2"] = loc.s["xl_l_tax_rate"]
    config["B2"] = float(TAX_RATE)

    # ── サマリシート（すべて数式。値は入れない）────────────
    d = loc.s["xl_sheet_detail"]
    c = loc.s["xl_sheet_config"]
    approved_label = loc.statuses[0]
    rows = (
        (loc.s["xl_l_subtotal"], f"=SUM('{d}'!D{FIRST_DATA_ROW}:D{last_row})"),
        (
            loc.s["xl_l_approved"],
            f"=SUMIF('{d}'!E{FIRST_DATA_ROW}:E{last_row},\"{approved_label}\","
            f"'{d}'!D{FIRST_DATA_ROW}:D{last_row})",
        ),
        (loc.s["xl_l_tax"], f"=ROUND(B2*'{c}'!B2,0)"),
        (loc.s["xl_l_total"], "=B2+B4"),
    )
    summary["A1"] = loc.fmt("xl_title", project=project, code=code)
    summary["A1"].font = Font(bold=True, size=13)
    for offset, (label, formula) in enumerate(rows):
        r = 2 + offset
        summary.cell(row=r, column=1, value=label).font = Font(bold=True)
        summary.cell(row=r, column=2, value=formula)
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 18

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    normalize_artifact(path)


def _build(rng: random.Random, outdir: Path, loc: Locale, code: str) -> tuple[dict[str, int], str]:
    project = rng.choice(loc.project_names)
    lines = _make_lines(rng, loc)
    rel = f"{SUBDIR}/{code}_quotation.xlsx"
    _write_workbook(outdir / rel, loc, code, project, lines)
    return _compute(lines, loc.statuses[0]), rel


def _unit(loc: Locale) -> str:
    return "円" if loc.code == "ja" else ""


def generate(
    rng: random.Random, outdir: Path, n_questions: int, locale: Locale
) -> list[Item]:
    """SPEC §3-3 の生成器契約（locale を追加したもの）。"""
    codes = [f"QT-{n:04d}" for n in rng.sample(range(1000, 5000), n_questions)]

    items: list[Item] = []
    for i, code in enumerate(codes):
        question_key, field, difficulty = QUESTION_SPECS[i % len(QUESTION_SPECS)]
        totals, rel = _build(rng, outdir, locale, code)
        value = totals[field]

        items.append(
            Item(
                qid=f"{CHANNEL}-{i + 1:02d}",
                channel=CHANNEL,
                question=locale.fmt(question_key, code=code),
                answer=str(value),
                answer_type="number",
                answer_aliases=number_aliases(value, _unit(locale))[1:],
                source_files=[rel],
                difficulty=difficulty,
                locale=locale.code,
            )
        )
    return items


def generate_fillers(
    rng: random.Random, outdir: Path, count: int, locale: Locale
) -> list[str]:
    """設問を持たない見積書を足す（ディストラクタ）。"""
    written: list[str] = []
    # 設問用コード（1000-4999）と番号帯を分けているので衝突しない
    for n in rng.sample(range(5000, 9999), count):
        _, rel = _build(rng, outdir, locale, f"QT-{n:04d}")
        written.append(rel)
    return written
