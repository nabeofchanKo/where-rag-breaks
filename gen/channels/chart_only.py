"""`chart_only` チャネル — 答えがグラフ画像の中の数値にしかない。

SPEC §3-1 #4 / §3-4。matplotlib で PNG を作り python-pptx で埋め込む。
**Office ファイルをレンダリングしないので LibreOffice は不要。**

なぜ古典的RAGが落ちるのか:
    python-pptx の図形テキスト抽出は、埋め込まれた画像の中身を見ない。
    OCR も掛けない（実務の既定の挙動）。したがって棒グラフの高さが表す数値は
    **抽出テキストに一切現れない**。

難易度 = 罠の機構:
    1 … **コントロール**。同じ値がスライド上の表にも書いてある。
    2 … グラフ画像にしかない。
    3 … グラフ画像にしかなく、本文には**全社合計**が書いてある。
        テキストだけ読むと、目に入る唯一の数値であるこれを答えてしまう。

値は 10 の倍数にしてある。棒の高さと目盛りから読める粒度にしておかないと、
画像が読めるアーム（Arm B の view_image）でも解けなくなり、
チャネルとして「誰も解けない」だけになってしまうため。
"""

from __future__ import annotations

import random
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from gen.common import Item, normalize_artifact
from gen.imaging import save_bar_chart
from gen.locales import Locale

CHANNEL = "chart_only"
SUBDIR = "reviews"
IMAGE_SUBDIR = "reviews/figures"

# (difficulty, 表にも値を出すか, 本文に全社合計を出すか)
QUESTION_SPECS: tuple[tuple[int, bool, bool], ...] = (
    (1, True, False),
    (2, False, False),
    (3, False, True),
)


def _build(
    rng: random.Random,
    outdir: Path,
    loc: Locale,
    code: str,
    show_table: bool,
    show_total: bool,
) -> tuple[str, int, int | None, list[str]]:
    sites = rng.sample(loc.sites, 4)
    # 10 の倍数。グラフから読み取れる粒度にする
    values = [rng.randrange(8, 60) * 10 for _ in sites]
    period = f"2027年度 第{rng.randrange(1, 5)}四半期" if loc.code == "ja" else "FY2027 Q2"
    asked_index = rng.randrange(len(sites))

    image_rel = f"{IMAGE_SUBDIR}/{code}_shipments.png"
    save_bar_chart(
        outdir / image_rel,
        loc.fmt("co_chart_title", period=period),
        list(sites),
        [float(v) for v in values],
        loc.s["co_chart_ylabel"],
        loc.code,
    )

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = loc.fmt("co_deck_title", period=period, code=code)

    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(4.2), Inches(3.4))
    frame = box.text_frame
    frame.word_wrap = True
    lines = [loc.s["co_body_caption"]]
    if show_table:
        lines.append(loc.s["co_table_header"])
        lines += [f"{site} | {value}" for site, value in zip(sites, values, strict=True)]
    if show_total:
        lines.append(loc.fmt("co_body_total", total=sum(values)))
    for index, line in enumerate(lines):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = line
        para.font.size = Pt(14)

    slide.shapes.add_picture(
        str(outdir / image_rel), Inches(4.9), Inches(1.4), width=Inches(4.6)
    )

    deck_rel = f"{SUBDIR}/{code}_review.pptx"
    path = outdir / deck_rel
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    normalize_artifact(path)

    decoy = sum(values) if show_total else None
    if decoy == values[asked_index]:
        decoy = None  # 偶然一致したら囮にならない
    return sites[asked_index], values[asked_index], decoy, [deck_rel, image_rel]


def generate(rng: random.Random, outdir: Path, n_questions: int, locale: Locale) -> list[Item]:
    codes = [f"CH-{n:04d}" for n in rng.sample(range(1000, 5000), n_questions)]

    items: list[Item] = []
    for i, code in enumerate(codes):
        difficulty, show_table, show_total = QUESTION_SPECS[i % len(QUESTION_SPECS)]
        site, value, decoy, files = _build(
            rng, outdir, locale, code, show_table, show_total
        )
        items.append(
            Item(
                qid=f"{CHANNEL}-{i + 1:02d}",
                channel=CHANNEL,
                question=locale.fmt("q_chart_only", code=code, site=site),
                answer=str(value),
                answer_type="number",
                answer_aliases=[f"{value:,}"] if value >= 1000 else [],
                source_files=files,
                decoys=[str(decoy)] if decoy is not None else [],
                difficulty=difficulty,
                locale=locale.code,
            )
        )
    return items


def generate_fillers(rng: random.Random, outdir: Path, count: int, locale: Locale) -> list[str]:
    written: list[str] = []
    for n in rng.sample(range(5000, 9999), count):
        _, show_table, show_total = QUESTION_SPECS[rng.randrange(len(QUESTION_SPECS))]
        _, _, _, files = _build(rng, outdir, locale, f"CH-{n:04d}", show_table, show_total)
        written.extend(files)
    return written
