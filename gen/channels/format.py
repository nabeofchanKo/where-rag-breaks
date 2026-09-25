"""`format` チャネル — 答えがセルの塗り色にある。

SPEC §3-1 #2。値そのものは手がかりにならない。

なぜ古典的RAGが落ちるのか:
    openpyxl でセルの値を読んでも、``cell.fill`` は取りに行かない限り出てこない。
    「実務で普通にやる範囲」の抽出（値をテキスト化する）では**塗り色は消える**。
    しかも消えたことに気づく手がかりが抽出テキスト側に残らない。

難易度 = 罠の機構（P0 の `formula` と同じ軸を通す）:
    1 … **コントロール**。塗られた行は備考欄にも「要再検査」と書いてある。
        テキストだけでも解ける段。ここが解けないなら抽出か検索の問題。
    2 … 塗り色だけ。備考欄は全行空。
    3 … 塗り色だけ。さらに**別の行**の備考欄に「確認中」と書いてある。
        テキストだけを読むと、目立つ文字列のあるこの行を答えてしまう。
        その行の管理番号を ``decoys`` に登録してある。
"""

from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from gen.common import Item, normalize_artifact
from gen.locales import Locale

CHANNEL = "format"
SUBDIR = "inspections"

HEADER_ROW = 3
FIRST_DATA_ROW = 4
YELLOW = PatternFill(start_color="FFFFF2A8", end_color="FFFFF2A8", fill_type="solid")

# (difficulty, 備考に印を出すか, 別行に囮の備考を置くか)
QUESTION_SPECS: tuple[tuple[int, bool, bool], ...] = (
    (1, True, False),
    (2, False, False),
    (3, False, True),
)


def _make_rows(rng: random.Random, loc: Locale) -> list[dict]:
    n = rng.randrange(9, 15)
    return [
        {
            "id": f"INS-{rng.randrange(10000, 99999)}",
            "item": loc.items[i % len(loc.items)],
            "value": round(rng.uniform(9.0, 12.0), 2),
            "spec": "10.0 ± 2.0",
        }
        for i in range(n)
    ]


def _write_workbook(
    path: Path,
    loc: Locale,
    code: str,
    site: str,
    rows: list[dict],
    flagged: int,
    note_on_flagged: bool,
    decoy: int | None,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = loc.s["fm_sheet"]

    ws["A1"] = loc.fmt("fm_title", site=site, code=code)
    ws["A1"].font = Font(bold=True, size=13)

    headers = ("fm_h_id", "fm_h_item", "fm_h_value", "fm_h_spec", "fm_h_note")
    for col, key in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=loc.s[key])
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for offset, row in enumerate(rows):
        r = FIRST_DATA_ROW + offset
        ws.cell(row=r, column=1, value=row["id"])
        ws.cell(row=r, column=2, value=row["item"])
        ws.cell(row=r, column=3, value=row["value"])
        ws.cell(row=r, column=4, value=row["spec"])

        if offset == flagged:
            # 答えはこの塗り色にしかない（難易度2・3の場合）
            for col in range(1, 6):
                ws.cell(row=r, column=col).fill = YELLOW
            if note_on_flagged:
                ws.cell(row=r, column=5, value=loc.s["fm_note_flagged"])
        elif offset == decoy:
            # 囮。テキストだけ読むとこの行が目立つ
            ws.cell(row=r, column=5, value=loc.s["fm_note_decoy"])

    for col, width in enumerate((14, 24, 12, 14, 20), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    normalize_artifact(path)


def _build(
    rng: random.Random, outdir: Path, loc: Locale, code: str, note: bool, with_decoy: bool
) -> tuple[str, str | None, str]:
    rows = _make_rows(rng, loc)
    flagged = rng.randrange(len(rows))
    decoy = None
    if with_decoy:
        others = [i for i in range(len(rows)) if i != flagged]
        decoy = rng.choice(others)

    rel = f"{SUBDIR}/{code}_inspection.xlsx"
    _write_workbook(
        outdir / rel, loc, code, rng.choice(loc.sites), rows, flagged, note, decoy
    )
    return rows[flagged]["id"], (rows[decoy]["id"] if decoy is not None else None), rel


def generate(rng: random.Random, outdir: Path, n_questions: int, locale: Locale) -> list[Item]:
    codes = [f"IL-{n:04d}" for n in rng.sample(range(1000, 5000), n_questions)]

    items: list[Item] = []
    for i, code in enumerate(codes):
        difficulty, note, with_decoy = QUESTION_SPECS[i % len(QUESTION_SPECS)]
        answer, decoy, rel = _build(rng, outdir, locale, code, note, with_decoy)
        items.append(
            Item(
                qid=f"{CHANNEL}-{i + 1:02d}",
                channel=CHANNEL,
                question=locale.fmt("q_format", code=code),
                answer=answer,
                answer_type="identifier",
                source_files=[rel],
                decoys=[decoy] if decoy and decoy != answer else [],
                difficulty=difficulty,
                locale=locale.code,
            )
        )
    return items


def generate_fillers(rng: random.Random, outdir: Path, count: int, locale: Locale) -> list[str]:
    written: list[str] = []
    for n in rng.sample(range(5000, 9999), count):
        _, note, with_decoy = QUESTION_SPECS[rng.randrange(len(QUESTION_SPECS))]
        _, _, rel = _build(rng, outdir, locale, f"IL-{n:04d}", note, with_decoy)
        written.append(rel)
    return written
