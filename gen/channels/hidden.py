"""`hidden` チャネル — 答えが発表者ノートにある。

SPEC §3-1 #10。既定を pptx の発表者ノートにする理由は §3-5 のとおりで、
python-docx にコメント追加 API が無く `word/comments.xml` の直接編集が要るため。

なぜ古典的RAGが落ちるのか:
    python-pptx で「図形のテキスト」を抽出するのが実務で普通にやる範囲だが、
    発表者ノートは本文の図形ツリーには入っていない（``slide.notes_slide``
    という別の場所にある）。**明示的に取りに行かない限り落ちる。**

難易度 = 罠の機構:
    1 … **コントロール**。同じ値がスライド本文にも書いてある。
    2 … ノートにしかない。
    3 … ノートにしかない。さらに本文に「前回案件の値引き率は{別の値}%」と
        書いてある。本文だけ読むとこちらを答える。``decoys`` に登録済み。
"""

from __future__ import annotations

import random
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from gen.common import Item, normalize_artifact
from gen.locales import Locale

CHANNEL = "hidden"
SUBDIR = "proposals"

# (difficulty, 本文にも答えを出すか, 本文に囮を置くか)
QUESTION_SPECS: tuple[tuple[int, bool, bool], ...] = (
    (1, True, False),
    (2, False, False),
    (3, False, True),
)


def _add_slide(prs: Presentation, title: str, body_lines: list[str], notes: str | None) -> None:
    layout = prs.slide_layouts[5]  # タイトルのみ。本文は自前のテキストボックスで置く
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title

    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(4.0))
    frame = box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(body_lines):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = line
        para.font.size = Pt(18)

    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _build(
    rng: random.Random,
    outdir: Path,
    loc: Locale,
    code: str,
    show_answer: bool,
    with_decoy: bool,
) -> tuple[int, int | None, str]:
    project = rng.choice(loc.project_names)
    site = rng.choice(loc.sites)
    rate = rng.randrange(5, 26)

    decoy_rate: int | None = None
    if with_decoy:
        # 正解と一致しない値にする（一致すると囮として機能しない）
        decoy_rate = rng.choice([r for r in range(3, 31) if r != rate])

    price_lines = [loc.s["hd_body_price"]]
    if show_answer:
        price_lines.append(loc.fmt("hd_body_discount_shown", rate=rate))
    if decoy_rate is not None:
        price_lines.append(loc.fmt("hd_body_discount_decoy", rate=decoy_rate))

    prs = Presentation()
    _add_slide(
        prs,
        loc.fmt("hd_title", project=project, code=code),
        [loc.fmt("hd_body_overview", site=site, project=project)],
        None,
    )
    _add_slide(
        prs,
        loc.s["hd_slide_price"],
        price_lines,
        # ★ 答えはここにしかない（難易度2・3）
        loc.fmt("hd_note_discount", rate=rate),
    )
    _add_slide(
        prs,
        loc.s["hd_slide_schedule"],
        [loc.fmt("hd_body_schedule", months=rng.randrange(3, 13))],
        None,
    )

    rel = f"{SUBDIR}/{code}_proposal.pptx"
    path = outdir / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    normalize_artifact(path)
    return rate, decoy_rate, rel


def generate(rng: random.Random, outdir: Path, n_questions: int, locale: Locale) -> list[Item]:
    codes = [f"PPT-{n:04d}" for n in rng.sample(range(1000, 5000), n_questions)]

    items: list[Item] = []
    for i, code in enumerate(codes):
        difficulty, show_answer, with_decoy = QUESTION_SPECS[i % len(QUESTION_SPECS)]
        rate, decoy_rate, rel = _build(rng, outdir, locale, code, show_answer, with_decoy)
        items.append(
            Item(
                qid=f"{CHANNEL}-{i + 1:02d}",
                channel=CHANNEL,
                question=locale.fmt("q_hidden", code=code),
                answer=str(rate),
                answer_type="number",
                answer_aliases=[f"{rate}%", f"{rate}％"],
                source_files=[rel],
                decoys=[str(decoy_rate)] if decoy_rate is not None else [],
                difficulty=difficulty,
                locale=locale.code,
            )
        )
    return items


def generate_fillers(rng: random.Random, outdir: Path, count: int, locale: Locale) -> list[str]:
    written: list[str] = []
    for n in rng.sample(range(5000, 9999), count):
        _, show_answer, with_decoy = QUESTION_SPECS[rng.randrange(len(QUESTION_SPECS))]
        _, _, rel = _build(rng, outdir, locale, f"PPT-{n:04d}", show_answer, with_decoy)
        written.append(rel)
    return written
