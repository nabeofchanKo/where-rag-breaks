"""`text` チャネル — 答えが本文の段落にある。

SPEC §3-1 #1。**このチャネルはベースライン**であり、古典的RAGが得意である
「べき」場所。H1（text では classical ≈ agentic）の検証対象なので、
意図的に難しくしてはいけない。ここで測っているのは情報チャネルの損失ではなく
**検索の質**である。

難易度は「同じ文に複数の候補が並ぶか」「素朴に読むと別の答えを拾うか」で
段階をつけている。いずれも 1 つのチャンクの中で答えられる。
"""

from __future__ import annotations

import datetime as dt
import random
from pathlib import Path

from docx import Document
from docx.shared import Pt

from gen.common import Item, normalize_artifact
from gen.locales import Locale

CHANNEL = "text"
SUBDIR = "projects"

# (質問キー, 答えのフィールド, answer_type, difficulty)
#   1 … 専用の文に1つだけ書いてある
#   2 … 1つの文に同種の候補が2つ並ぶ（窓口と技術照会 / 着手日と完了日）
#   3 … 主管部署と決裁部署が別。素朴に読むと主管部署を答えてしまう
QUESTION_SPECS: tuple[tuple[str, str, str, int], ...] = (
    ("q_owner", "owner", "name", 1),
    ("q_site", "site", "name", 1),
    ("q_department", "department", "name", 1),
    ("q_engineer", "engineer", "name", 2),
    ("q_end", "end", "text", 2),
    ("q_approver", "approver", "name", 3),
)


def _make_facts(rng: random.Random, loc: Locale, code: str) -> dict[str, str]:
    """1 文書ぶんの事実を決める。ここで確定した値がそのまま正解になる。"""
    owner, engineer = rng.sample(loc.people, 2)
    department, approver = rng.sample(loc.departments, 2)  # 必ず別部署（difficulty 3 の肝）
    start = dt.date(2027, 1, 1) + dt.timedelta(days=rng.randrange(0, 200))
    end = start + dt.timedelta(days=rng.randrange(60, 300))
    return {
        "code": code,
        "site": rng.choice(loc.sites),
        "project": rng.choice(loc.project_names),
        "owner": owner,
        "engineer": engineer,
        "department": department,
        "approver": approver,
        "start": _format_date(start, loc),
        "end": _format_date(end, loc),
        "start_iso": start.isoformat(),
        "end_iso": end.isoformat(),
    }


def _format_date(d: dt.date, loc: Locale) -> str:
    """文書中の表示形式。設問は ISO 形式を要求し、表示形式は alias で受理する。"""
    if loc.code == "ja":
        return f"{d.year}年{d.month}月{d.day}日"
    return d.strftime("%d %B %Y")


def _date_aliases(iso: str, loc: Locale) -> list[str]:
    y, m, d = (int(x) for x in iso.split("-"))
    date = dt.date(y, m, d)
    aliases = [iso, f"{y}/{m:02d}/{d:02d}", f"{y}/{m}/{d}"]
    if loc.code == "ja":
        aliases += [f"{y}年{m}月{d}日", f"{y}年{m:02d}月{d:02d}日"]
    else:
        aliases += [date.strftime("%d %B %Y"), date.strftime("%B %d, %Y")]
    return aliases


def _write_document(path: Path, loc: Locale, facts: dict[str, str]) -> None:
    """案件概要書の .docx を書く。設問の有無で本文は変えない（罠を仕込まない）。"""
    doc = Document()
    style = doc.styles["Normal"].font
    style.name = "Yu Gothic" if loc.code == "ja" else "Calibri"
    style.size = Pt(10.5)

    doc.add_heading(
        loc.fmt("doc_header", title=loc.s["doc_title"], code=facts["code"]), level=1
    )

    sections = (
        ("sec_outline", "outline_body"),
        ("sec_scope", "scope_body"),
        ("sec_schedule", "schedule_body"),
        ("sec_contact", "contact_body"),
        ("sec_notes", "notes_body"),
    )
    for heading_key, body_key in sections:
        doc.add_heading(loc.s[heading_key], level=2)
        doc.add_paragraph(loc.fmt(body_key, **facts))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    normalize_artifact(path)


def _build(
    rng: random.Random, outdir: Path, loc: Locale, code: str, spec_index: int | None
) -> tuple[dict[str, str], str]:
    """1 文書を生成して (事実, コーパス相対パス) を返す。"""
    facts = _make_facts(rng, loc, code)
    rel = f"{SUBDIR}/{code}_overview.docx"
    _write_document(outdir / rel, loc, facts)
    return facts, rel


def generate(
    rng: random.Random, outdir: Path, n_questions: int, locale: Locale
) -> list[Item]:
    """SPEC §3-3 の生成器契約（locale を追加したもの）。

    ``outdir`` は corpus/files。返す ``source_files`` は corpus/files からの相対パス。
    """
    # 案件コードは重複させない。質問文の鍵になるため。
    codes = [f"PRJ-{n:04d}" for n in rng.sample(range(1000, 5000), n_questions)]

    items: list[Item] = []
    for i, code in enumerate(codes):
        question_key, field, answer_type, difficulty = QUESTION_SPECS[i % len(QUESTION_SPECS)]
        facts, rel = _build(rng, outdir, locale, code, i)

        if field == "end":
            answer = facts["end_iso"]
            aliases = _date_aliases(facts["end_iso"], locale)
        else:
            answer = facts[field]
            # 日本語の氏名は文書中では姓名の間に空白がある。空白なしも受理する。
            aliases = [answer.replace(" ", "")] if " " in answer else []
            if field == "site":
                # 作業範囲の節が「作業範囲は{site}構内に限る」と書いているため、
                # モデルは「{site}構内」と答える。場所としては同一なので受理する。
                # ★ これは結果を見たあとに追加した alias である。
                #   経緯と、追加前後の run を比較してはいけない旨は
                #   docs/scoring-changes.md に記載してある（SPEC §11）。
                aliases.append(locale.fmt("site_premises", site=answer))

        items.append(
            Item(
                qid=f"{CHANNEL}-{i + 1:02d}",
                channel=CHANNEL,
                question=locale.fmt(question_key, code=code),
                answer=answer,
                answer_type=answer_type,
                answer_aliases=aliases,
                source_files=[rel],
                difficulty=difficulty,
                locale=locale.code,
            )
        )
    return items


def generate_fillers(
    rng: random.Random, outdir: Path, count: int, locale: Locale
) -> list[str]:
    """設問を持たない案件概要書を足す。

    検索を実際に働かせるためのディストラクタ。設問つき文書と**見分けがつかない
    形式**であることが重要（形式で当てられてしまうと検索の評価にならない）。
    """
    written: list[str] = []
    # 設問用コード（1000-4999）と番号帯を分けているので衝突しない
    for n in rng.sample(range(5000, 9999), count):
        code = f"PRJ-{n:04d}"
        _, rel = _build(rng, outdir, locale, code, None)
        written.append(rel)
    return written
