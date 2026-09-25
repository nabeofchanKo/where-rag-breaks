"""Arm A のテキスト抽出。

★ **SPEC §4-1 の大原則**: 抽出器に意図的な穴を作らない。
古典的RAGが落ちるのは**構造的な理由**であるべきで、実装の手抜きであっては
ならない。したがってここでは「実務で普通にやる範囲」を素直に、かつ丁寧に実装する。

具体的に、手を抜かないとはこういうこと:
    - xlsx は値だけでなく**数式文字列も出す**。data_only=True だけで読むと
      数式セルが全部 None になり、それは抽出器の穴になってしまう
    - xlsx はセル番地を付けて**行の対応関係を保ったまま**出す。列を潰して
      平文にすると、それは抽出器の穴になってしまう
    - docx は見出し階層を保ち、表もセル区切りを保って出す

そのうえで落ちるなら、それがこのベンチマークの主張である。

★ **抽出範囲の線引き（開示事項）**
    SPEC §4-1 は抽出範囲を「python-docx のテキスト、openpyxl の値、
    pypdf/pdfplumber のテキスト、python-pptx の**図形テキスト**」と定めている。
    この線引きの結果、次のものは抽出されない。

    - **pptx の発表者ノート** … スライドの図形ツリーではなく
      ``slide.notes_slide`` という別の場所にある。`hidden` チャネルの前提そのもの
    - **セルの塗り色・文字色・太字** … 値の抽出には現れない。`format` の前提
    - **画像の中身** … OCR を掛けない。`chart_only` / `layout` / `scanned` の前提

    これらを取りに行けば該当チャネルは定義上消える。取りに行かないのは
    「実務の既定の挙動を測る」というこのベンチマークの目的によるもので、
    隠したい弱点ではない。**ここに明記したうえで測っている。**
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# 抽出できる拡張子。ここに無いものは「本文テキストが取れないファイル」として扱う
# （PNG など。OCR は掛けない）。
SUPPORTED = {".docx", ".xlsx", ".pptx", ".pdf"}


@dataclass(frozen=True)
class Block:
    """抽出された構造単位。チャンク分割の入力になる。

    ``locator`` は根拠の所在（見出し名 / シート名と行範囲）。evidence の
    突き合わせと、人間が結果を追うために持つ。
    """

    path: str  # コーパス相対パス
    locator: str
    text: str

    @property
    def header(self) -> str:
        return f"[{self.path} :: {self.locator}]"


# ── docx ────────────────────────────────────────────────────────────
def _iter_body(doc: Document) -> list[Paragraph | Table]:
    """本文の段落と表を、文書に現れる順序どおりに返す。"""
    from docx.oxml.ns import qn

    out: list[Paragraph | Table] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            out.append(Paragraph(child, doc))
        elif child.tag == qn("w:tbl"):
            out.append(Table(child, doc))
    return out


def _table_to_text(table: Table) -> str:
    """表はセル区切りを残して出す（列の対応を潰さない）。"""
    rows = []
    for row in table.rows:
        rows.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(rows)


def extract_docx(path: Path, rel: str) -> list[Block]:
    """見出しごとにまとめる。見出し階層が文書の構造そのものなので保持する。"""
    doc = Document(path)
    blocks: list[Block] = []
    current_heading = "(冒頭)"
    buffer: list[str] = []

    def flush() -> None:
        body = "\n".join(t for t in buffer if t.strip())
        if body.strip():
            blocks.append(Block(path=rel, locator=current_heading, text=body))
        buffer.clear()

    for element in _iter_body(doc):
        if isinstance(element, Table):
            buffer.append(_table_to_text(element))
            continue

        style = (element.style.name or "") if element.style is not None else ""
        text = element.text.strip()
        if style.startswith("Heading") or style.startswith("Title"):
            flush()
            current_heading = text or current_heading
            buffer.append(text)
        elif text:
            buffer.append(text)

    flush()
    return blocks


# ── xlsx ────────────────────────────────────────────────────────────
def _cell_repr(value: object, formula: object) -> str | None:
    """1 セルの表示。値と数式の**両方**を出す。

    openpyxl が書いた数式セルにはキャッシュ値が無く、``data_only=True`` では
    None になる。片方しか出さないのは抽出器の穴なので、両方出す。
    """
    has_value = value is not None and str(value).strip() != ""
    is_formula = isinstance(formula, str) and formula.startswith("=")

    if is_formula and has_value:
        return f"{value} (式: {formula})"
    if is_formula:
        return str(formula)
    if has_value:
        return str(value)
    return None


def extract_xlsx(path: Path, rel: str) -> list[Block]:
    """シートごとに、セル番地つきで行の対応を保ったまま出す。"""
    wb_values = load_workbook(path, data_only=True, read_only=True)
    wb_formulas = load_workbook(path, data_only=False, read_only=True)

    blocks: list[Block] = []
    try:
        for sheet_name in wb_formulas.sheetnames:
            ws_v = wb_values[sheet_name]
            ws_f = wb_formulas[sheet_name]

            lines: list[str] = []
            for row_v, row_f in zip(ws_v.iter_rows(), ws_f.iter_rows(), strict=False):
                cells = []
                for cell_v, cell_f in zip(row_v, row_f, strict=False):
                    shown = _cell_repr(cell_v.value, cell_f.value)
                    if shown is not None:
                        ref = f"{get_column_letter(cell_f.column)}{cell_f.row}"
                        cells.append(f"{ref}: {shown}")
                if cells:
                    lines.append(" | ".join(cells))

            if lines:
                blocks.append(
                    Block(path=rel, locator=f"シート {sheet_name}", text="\n".join(lines))
                )
    finally:
        wb_values.close()
        wb_formulas.close()

    return blocks


# ── pptx ────────────────────────────────────────────────────────────
def _shape_text(shape) -> str:
    """図形のテキスト。表は行ごとにセル区切りを残す。"""
    if shape.has_table:
        return "\n".join(
            " | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows
        )
    if shape.has_text_frame:
        return "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
    return ""


def extract_pptx(path: Path, rel: str) -> list[Block]:
    """スライドごとに図形テキストをまとめる。

    ★ 発表者ノート（``slide.notes_slide``）は**取らない**。SPEC §4-1 の
    「python-pptx の図形テキスト」という線引きによるもので、`hidden` チャネルの
    前提そのもの。モジュール冒頭の開示事項を参照。
    """
    from pptx import Presentation

    prs = Presentation(path)
    blocks: list[Block] = []
    for index, slide in enumerate(prs.slides, start=1):
        parts = [t for t in (_shape_text(shape) for shape in slide.shapes) if t.strip()]
        if parts:
            blocks.append(
                Block(path=rel, locator=f"スライド {index}", text="\n".join(parts))
            )
    return blocks


# ── pdf ─────────────────────────────────────────────────────────────
def extract_pdf(path: Path, rel: str) -> list[Block]:
    """ページごとにテキストを取る。

    pdfplumber を先に試し、取れなければ pypdf に落とす（実務で普通にやる範囲）。
    画像だけの PDF（`scanned` チャネル）はどちらでも空になる。これは抽出器の
    穴ではなく、**テキスト層が存在しない**という構造。
    """
    texts: list[str] = []
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            texts = [(page.extract_text() or "") for page in pdf.pages]
    except Exception:  # noqa: BLE001 — pdfplumber が読めない PDF は pypdf に任せる
        texts = []

    if not any(t.strip() for t in texts):
        try:
            from pypdf import PdfReader

            texts = [(page.extract_text() or "") for page in PdfReader(str(path)).pages]
        except Exception:  # noqa: BLE001
            texts = []

    return [
        Block(path=rel, locator=f"ページ {i}", text=text)
        for i, text in enumerate(texts, start=1)
        if text.strip()
    ]


# ── 入口 ────────────────────────────────────────────────────────────
_EXTRACTORS = {
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".pptx": extract_pptx,
    ".pdf": extract_pdf,
}


def extract_file(path: Path, rel: str) -> list[Block]:
    extractor = _EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        return []
    return extractor(path, rel)


def extract_corpus(files_dir: Path) -> list[Block]:
    """corpus/files 配下を全部抽出する。順序はパス昇順で決定的。"""
    blocks: list[Block] = []
    for path in sorted(files_dir.rglob("*")):
        if path.is_file():
            blocks.extend(extract_file(path, path.relative_to(files_dir).as_posix()))
    return blocks
